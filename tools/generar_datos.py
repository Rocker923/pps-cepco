#!/usr/bin/env python3
"""
Genera data/actividades.js a partir de las bitácoras semanales (.xlsx).

Uso (desde la raíz del repositorio):
    pip install openpyxl
    python tools/generar_datos.py                # lee ./bitacoras/*.xlsx
    python tools/generar_datos.py ruta/a/carpeta # lee otra carpeta

Reglas de lectura (Hoja1 de cada archivo):
  - Encabezado: E6 = # de semana, G6 = fecha inicio, I6 = fecha fin.
  - Actividades: desde la fila 8 hacia abajo (A = número, B = actividad,
    F = día de ejecución, G = observaciones) hasta la fila de "horas".
  - Horas: fila "Numero de horas acumuladas totales" (columna G).
  - Desglose: filas posteriores al título "Desglose de actividades".
    Un renglón con número en A inicia un ítem; los renglones sin número
    son continuación del ítem anterior.

Las semanas cuyo desglose se entregó escaneado como imagen (1, 2 y 3) y el
ítem 18 de la semana 4 se transcriben en DESGLOSE_MANUAL. Si se corrige el
texto en el Excel, basta volver a ejecutar este script.
"""
import datetime as dt
import difflib
import glob
import json
import os
import re
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instale con:  pip install openpyxl")

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, "data", "actividades.js")
HORAS_META = 800

# ---------------------------------------------------------------------------
# Desgloses transcritos desde las imágenes escaneadas insertadas en el Excel.
# Clave: número de semana -> {número de actividad: texto}
# ---------------------------------------------------------------------------
DESGLOSE_MANUAL = {
    1: {
        1: "Toma de datos personales, contactos de emergencia, referencias personales. Charla obligatoria de ciberseguridad y evaluación de conocimiento en el tema. Charla obligatoria de principios y valores de la empresa. Recorrido preliminar por todos los sistemas principales de la planta, desde el alimentador de carbón, las bandas transportadoras, molino de carbón, bunker, alimentadores de carbón, spreader, rejilla movible, escotillas para ver el interior de la caldera, sistema hopper, electrostatic precipitator, ID fan, sistema de aire comprimido de instrumentación y de servicio, condensador, turbina, reductor, hasta el generador.",
        2: "Explicación de diagramas sobre las líneas de aceite para lubricación del arreglo turbina-reductor-generador, niveles de aceite aceptables, presiones aceptables de operación y de standby, sistema de retorno al tanque de reserva.",
        3: "Explicación de diagrama unifilar del sistema de recepción y distribución eléctrica principal de la planta. Explicación detallada de datos sobre condiciones estables de funcionamiento y lectura de diagrama de ciclo del agua para las Cooling Towers.",
        4: "Reunión con el jefe de mantenimiento mecánico para la explicación de propiedades específicas y valores del sistema de lubricación y torres de enfriamiento. Lectura de análisis realizado a la calidad del aceite actual de la turbina-reductor-generador, descripción de los rodamientos y su lubricación especial. Investigación para la interpretación de valores en el reporte del estado actual del aceite de lubricación.",
        5: "Descripción de la ruta del sistema de gestión para generar los \"pedidos\" y solicitar los productos en las bodegas, también para generar órdenes de trabajo. Observar y cumplir con la seguridad industrial para el proceso de elevación de un tubo de aleación desde nivel cero hasta el nivel 6 de la caldera. Observar el ajuste fino del accionador de un sistema steam blower.",
        6: "Conversación técnica sobre eficiencia de consumo de carbón, composición química del carbón y reconocimiento de equipo de alineación por láser. Reconocimiento de equipo para termografía, ultrasonido de espesores.",
        7: "Planteamiento de posibles proyectos a desarrollar durante la práctica profesional: 1) Diseño, montaje y mantenimiento de una esclusa para el silo de ceniza húmeda. 2) Diseño, montaje y mantenimiento de compuertas para pesaje de ceniza húmeda. 3) Desarrollo e implementación de software interactivo y amigable para el cálculo de la eficiencia de la caldera. 4) Levantamiento de inventario y fichas técnicas de los rodamientos de toda la planta.",
        8: "Lectura de planos y manual de mantenimiento para sootblowers, su funcionamiento y utilidad. Investigación sobre tipos de válvulas usadas en la planta. Revisión de checklists para órdenes de trabajo. Evaluación de factores para cambio de lubricación en los spreaders del stoker. Investigación sobre Forced Draft Fan, Induced Draft Fan, Secondary Air System, Electrostatic Precipitator.",
        9: "Observar la reparación de la banda sumergida, tipos de materiales y pegamentos especiales, cambio de rodamiento de motor FD Fan, investigación in situ para cambio y soldadura de válvula de venteo de la main steam. Datos importantes sobre aleaciones de tubos y válvulas y los tipos de soldaduras apropiadas. Investigación sobre bomba multietapa.",
        10: "Investigación de campo sobre el tipo y características de la caldera, de la rejilla movible, de la entrada de aire precalentado. Solicitud y entrega de equipo para alineación de un sistema motor-acople elástico-caja reductora, y luego puesta a punto del sistema de medición láser. Toma de puntos de referencia para iniciar el proceso de alineación, nivelación por shims y la alineación por medio de tornillos de ajuste soldados a la bancada.",
    },
    2: {
        1: "Observar la desactivación del sistema de accionamiento neumático. Presenciar el desmontaje del sleeve de compensación. Retirar la tornillería del empaque de la válvula domo. Inspección del estado de la válvula domo. Retirar el sello dañado.",
        2: "Escuchar y ver atentamente la explicación de los sistemas de ventiladores de aire que van hacia la caldera y los que extraen gases de escape.",
        3: "Investigación bibliográfica sobre tipos de acoples entre motores eléctricos y otros equipos usados en la planta. Lectura e interpretación de planos de bombas.",
        4: "Observar y documentar el proceso completo de la elaboración técnica/artesanal de un sello elástico para la unión de tuberías de lubricación.",
        5: "Consultar al personal indicado sobre el manejo responsable de desechos potencialmente contaminantes y las políticas estrictas de la institución para la activa conservación del medio ambiente. Los ingenieros de seguridad industrial han sido muy amables en explicarme a detalle los procedimientos.",
        6: "Investigación sobre modos de falla y soluciones a problemas relacionados con sistemas de bombeo hidráulico.",
        7: "Investigación en campo sobre la frecuencia y tipos de mantenimiento realizados a la bomba y filtros de aceite de lubricación y de control.",
        8: "Recopilación de los mantenimientos realizados a los equipos de la planta durante el mes de julio.",
        9: "Inicio del levantamiento del inventario para la realización de fichas técnicas y frecuencia de inspección para los rodamientos en los equipos de la planta.",
    },
    3: {
        1: "Organizar archivos físicos de órdenes de trabajo en fechas ascendentes. Revisar en las órdenes de trabajo las observaciones de equipos que presentan fallas.",
        2: "Como parte de la investigación para el desarrollo del software para cálculo de la caldera realicé una investigación bibliográfica sobre contaminantes NOx y SOx.",
        3: "Investigación en campo sobre sistemas de precalentamiento de aire, su funcionamiento y utilidad. Los operarios han sido muy amables en explicarme.",
        4: "Inicio de lectura de la norma ASME PTC 4.1 para cálculo de la eficiencia de la caldera. Fórmulas específicas, cálculo de factores de corrección.",
        5: "Explicación y búsqueda de datos en campo sobre temperaturas y presiones del domo. Los técnicos e ingenieros del sistema DCS han sido muy amables en explicarme a detalle los datos recolectados.",
        6: "Investigación bibliográfica sobre parámetros GCV, NCV. Datos sobre el combustible que influyen en la eficiencia de la caldera y el consumo de carbón.",
        7: "Observar los sistemas para quitar hollín en funcionamiento cuando la planta está en línea. Determinar si se comportan de forma normal o anormal.",
        8: "Planteamiento de proyecto de recirculación de ceniza. Observaciones y deliberaciones sobre espacio disponible en la planta para instalar el recirculador.",
        9: "Investigación bibliográfica y comprobación en campo sobre seguros retenedores. Norma DIN/ASME B18.27.",
        10: "Inspección y realización de termografía en motor de torre de enfriamiento. Capacitación in situ para realizar una termografía correctamente. Toma de imágenes termográficas e interpretación de posibles modos de falla para el arreglo motor – transmisión – caja reductora de una torre de enfriamiento.",
    },
    4: {
        18: "Organizar archivo de órdenes de trabajo por fecha y leer las observaciones que los operarios hacen de los equipos en mantenimiento.",
    },
}

# ---------------------------------------------------------------------------
# Áreas de trabajo (colores del Gantt). Se asignan por palabras clave en el
# orden indicado; la primera que coincida gana. Edite libremente.
# ---------------------------------------------------------------------------
AREAS = [
    ("turbina", "Nueva turbina 1.6 MW", [
        "nueva turbina", "1.6 mw", "sala de máquinas", "sala de maquinas", "plano de planta",
        "bloque monolítico", "cimentación", "isométric", "tubería de vapor", "tuberia de vapor",
        "boquillas", "solidworks", "fusion 360", "fuision 360", "renderizado", "mallado", "mallas",
        "obj", "sldprt", "pdf a dxf", "dxf", "autocad", "anexo", "pórtico", "grúa", "mesh",
        "dibujo técnico", "dibujo 3d", "dibujo de planta", "superficie 3d", "nwd", "planos realizados",
        "visualización", "boceto", "escalar plano", "aplanado de líneas", "sólidos",
        "generación gráfica", "dimensionamiento de la nueva", "ruta de tubería", "planos autocad",
        "unión entre la tubería",
    ]),
    ("agni", "Software Agni · eficiencia de caldera", [
        "agni", "software para cálculo", "software de cálculo", "software para diseño",
        "asme ptc", "gcv", "ncv", "nox", "pérdidas a considerar", "consumo específico",
        "implementación de software", "entorno de desarrollo",
    ]),
    ("sap", "Actualización SAP e inventario", [
        "sap", "plantilla", "inventario", "rodamiento", "número de series", "numero de series",
        "datos de equipos", "datos de varios equipos", "base de datos", "grasa usada",
    ]),
    ("docs", "Gestión de órdenes de trabajo", [
        "órdenes de trabajo", "ordenes de trabajo", "mantenimiento general realizado",
    ]),
    ("formacion", "Inducción, formación y vinculación", [
        "expediente", "charla", "reunión con", "explicación", "explicacion", "investigación bibliográfica",
        "investigacion bibliográfica", "investigacipon", "estudio de normativa", "estudio sobre",
        "lectura de", "conversasión", "plática", "planteamiento de posibles", "pensum",
        "estudiante", "aprender", "ruta didáctica", "incapacidad", "deliberación",
        "investigación sobre", "discusión", "planteo",
    ]),
    ("campo", "Mantenimiento e inspección en campo", []),  # valor por defecto
]


def norm(t):
    t = unicodedata.normalize("NFKD", (t or "").lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", t)).strip()


def area_de(texto):
    t = (texto or "").lower()
    for clave, _nombre, palabras in AREAS:
        if any(p in t for p in palabras):
            return clave
    return AREAS[-1][0]


def limpiar(t):
    if t is None:
        return ""
    return re.sub(r"\s+", " ", str(t)).strip()


def parse_fecha(v, ini, fin, avisos, ref):
    """Devuelve date dentro de la semana; corrige errores de captura comunes."""
    d = None
    if isinstance(v, dt.datetime):
        d = v.date()
    elif isinstance(v, dt.date):
        d = v
    elif isinstance(v, str):
        s = v.strip()
        m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})$", s)          # 14/9/206
        m2 = re.match(r"^(\d{1,2})/(\d{1,2})(\d{4})$", s)              # 10/82026
        if m:
            dd, mm = int(m.group(1)), int(m.group(2))
        elif m2:
            dd, mm = int(m2.group(1)), int(m2.group(2))
        else:
            dd = mm = None
        if dd:
            try:
                d = dt.date(ini.year, mm, dd)
                avisos.append(f"{ref}: fecha escrita como texto '{s}' → {d.isoformat()}")
            except ValueError:
                d = None
    if d is None:
        return None
    if not (ini <= d <= fin):
        for mes in (ini.month, fin.month):
            try:
                c = d.replace(year=ini.year, month=mes)
            except ValueError:
                continue
            if ini <= c <= fin:
                avisos.append(f"{ref}: fecha {d.isoformat()} fuera de la semana → corregida a {c.isoformat()}")
                return c
        avisos.append(f"{ref}: fecha {d.isoformat()} fuera de la semana ({ini}–{fin}); se deja tal cual")
    return d


def sim(titulo, desglose):
    a, b = norm(titulo), norm(desglose)
    if not a or not b:
        return 0.0
    return difflib.SequenceMatcher(None, a, b[: len(a) + 25]).ratio()


def alinear(acts, items):
    """Alineación monótona (programación dinámica) actividad ↔ ítem de desglose."""
    n, m = len(acts), len(items)
    S = [[sim(a["titulo"], it) for it in items] for a in acts]
    INF = -1e9
    best = [[INF] * (m + 1) for _ in range(n + 1)]
    back = [[None] * (m + 1) for _ in range(n + 1)]
    best[0] = [0.0] * (m + 1)
    for i in range(n + 1):
        best[i][0] = 0.0
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            opts = [
                (best[i - 1][j - 1] + S[i - 1][j - 1] - 0.25, "d"),
                (best[i - 1][j], "u"),
                (best[i][j - 1], "l"),
            ]
            best[i][j], back[i][j] = max(opts)
    i, j, par = n, m, {}
    while i > 0 and j > 0:
        mv = back[i][j]
        if mv == "d":
            if S[i - 1][j - 1] >= 0.35:
                par[i - 1] = j - 1
            i, j = i - 1, j - 1
        elif mv == "u":
            i -= 1
        else:
            j -= 1
    # Relleno de huecos: si entre dos pares consecutivos queda exactamente una
    # actividad y un ítem sin pareja, se emparejan (desgloses que no repiten el título).
    usados = set(par.values())
    anclas = [(-1, -1)] + sorted(par.items()) + [(n, m)]
    for (i0, j0), (i1, j1) in zip(anclas, anclas[1:]):
        ai = [i for i in range(i0 + 1, i1) if i not in par]
        bj = [j for j in range(j0 + 1, j1) if j not in usados]
        if len(ai) == 1 and len(bj) == 1:
            par[ai[0]] = bj[0]
    return par


def leer_semana(ruta, desglose_previo, avisos):
    ws = openpyxl.load_workbook(ruta, data_only=True).worksheets[0]
    semana = int(ws["E6"].value)
    ini, fin = ws["G6"].value.date(), ws["I6"].value.date()
    horas_sem = horas_tot = None
    acts, items = [], []
    modo = "act"
    for r in range(8, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        d_, e_ = ws.cell(r, 4).value, ws.cell(r, 5).value
        if isinstance(a, str) and a.strip().upper().startswith("AL ENVIAR"):
            break
        if isinstance(e_, str) and "desglose" in e_.lower():
            modo = "des"
            continue
        if isinstance(d_, str) and "horas" in d_.lower():
            g = ws.cell(r, 7).value
            if "totales" in d_.lower():
                horas_tot = g
            else:
                horas_sem = g
            continue
        if isinstance(e_, str) and ("Lic." in e_ or "Gerente" in e_):
            continue
        num = None
        if a is not None:
            mnum = re.match(r"^\s*(\d+)", str(a))
            num = int(mnum.group(1)) if mnum else None
        if modo == "act":
            titulo = limpiar(b)
            if num and titulo:
                ref = f"Semana {semana}, actividad {num}"
                acts.append({
                    "n": num,
                    "titulo": titulo,
                    "fecha": parse_fecha(ws.cell(r, 6).value, ini, fin, avisos, ref),
                    "obs": limpiar(ws.cell(r, 7).value),
                })
        else:
            txt = limpiar(b)
            if num:
                items.append(txt)
            elif txt and items:
                items[-1] = (items[-1] + " " + txt).strip()

    # --- desglose -----------------------------------------------------------
    items = [t for t in items if t]
    manual = DESGLOSE_MANUAL.get(semana, {})
    desg = {}
    if items and desglose_previo and [norm(x) for x in items] == [norm(x) for x in desglose_previo]:
        avisos.append(f"Semana {semana}: el desglose es copia idéntica del de la semana anterior; se ignora.")
        items = []
    if items:
        # Actividades ya cubiertas por transcripción manual no compiten por ítems
        idx_libres = [i for i, a in enumerate(acts) if a["n"] not in manual]
        if len(items) == len(idx_libres):
            par = {i: k for k, i in enumerate(idx_libres)}
            todos = [norm(x) for x in items] + [norm(x) for x in (desglose_previo or [])]
            for i in list(par):
                txt = items[par[i]]
                propio = norm(txt)
                dup = sum(1 for t in todos if t == propio) > 1
                if dup and sim(acts[i]["titulo"], txt) < 0.5:
                    avisos.append(f"Semana {semana}, actividad {acts[i]['n']}: el desglose repite el de otra "
                                  f"actividad («{txt[:50]}…»); se omite.")
                    del par[i]
        else:
            par = alinear(acts, items)
            sin = [acts[i]["n"] for i in range(len(acts)) if i not in par]
            if sin:
                avisos.append(f"Semana {semana}: {len(acts)} actividades y {len(items)} ítems de desglose; "
                              f"sin desglose las actividades {sin}.")
        for i, j in par.items():
            desg[acts[i]["n"]] = items[j]
    for n, t in manual.items():
        desg[n] = t

    filas = []
    for a in acts:
        aid = f"S{semana:02d}-A{a['n']:02d}"
        texto_area = a["titulo"] + " " + desg.get(a["n"], "")
        filas.append({
            "id": aid,
            "semana": semana,
            "n": a["n"],
            "titulo": a["titulo"],
            "fecha": a["fecha"].isoformat() if a["fecha"] else None,
            "observacion": a["obs"],
            "desglose": desg.get(a["n"], ""),
            "area": area_de(a["titulo"]) if area_de(a["titulo"]) != "campo" else area_de(texto_area),
        })
    return {
        "semana": semana,
        "inicio": ini.isoformat(),
        "fin": fin.isoformat(),
        "horasSemana": horas_sem,
        "horasAcumuladas": horas_tot,
        "actividades": filas,
    }, items


def main():
    carpeta = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, "bitacoras")
    rutas = glob.glob(os.path.join(carpeta, "*.xlsx"))
    rutas = [r for r in rutas if not os.path.basename(r).startswith("~$")]
    if not rutas:
        sys.exit(f"No se encontraron .xlsx en {carpeta}")

    def clave(r):
        m = re.search(r"(\d+)", os.path.basename(r))
        return int(m.group(1)) if m else 0

    avisos, semanas, previo = [], [], None
    for r in sorted(rutas, key=clave):
        s, previo = leer_semana(r, previo, avisos)
        semanas.append(s)
    semanas.sort(key=lambda s: s["semana"])

    acumulado = 0
    for s in semanas:
        acumulado += s["horasSemana"] or 0
        if s["horasAcumuladas"] is not None and s["horasAcumuladas"] != acumulado:
            avisos.append(f"Semana {s['semana']}: horas acumuladas en el Excel = {s['horasAcumuladas']}, "
                          f"pero la suma semanal da {acumulado}; la página usa {acumulado}.")
        s["horasAcumuladas"] = acumulado

    datos = {
        "generado": dt.datetime.now().isoformat(timespec="seconds"),
        "horasMeta": HORAS_META,
        "areas": [{"id": k, "nombre": n} for k, n, _ in AREAS],
        "semanas": semanas,
    }
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    with open(SALIDA, "w", encoding="utf-8") as f:
        f.write("/* Archivo generado por tools/generar_datos.py — no editar a mano. */\n")
        f.write("window.PPS_DATOS = ")
        json.dump(datos, f, ensure_ascii=False, indent=1)
        f.write(";\n")

    # Lista de códigos para saber en qué carpeta de media/ va cada archivo
    lista = os.path.join(RAIZ, "media", "CODIGOS.md")
    with open(lista, "w", encoding="utf-8") as f:
        f.write("# Códigos de actividad\n\n")
        f.write("Cree una carpeta `media/<CÓDIGO>/` y coloque ahí fotos, planos (PDF, DWG, DXF…) y, si desea, un `notas.md`.\n\n")
        for s in semanas:
            f.write(f"\n## Semana {s['semana']} ({s['inicio']} a {s['fin']})\n\n")
            for a in s["actividades"]:
                f.write(f"- `{a['id']}` — {a['fecha'] or 's/f'} — {a['titulo']}\n")

    total = sum(len(s["actividades"]) for s in semanas)
    print(f"OK: {len(semanas)} semanas, {total} actividades → {os.path.relpath(SALIDA, RAIZ)}")
    if avisos:
        print("\nAvisos (revise en el Excel si desea corregir la fuente):")
        for a in avisos:
            print("  -", a)


if __name__ == "__main__":
    main()
